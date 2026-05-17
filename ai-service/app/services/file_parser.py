"""
File Parser — Extract text from PDF, DOCX, and plain text files.
Downloads files from URLs and parses them using pdfplumber and python-docx.
"""

import io
from typing import Optional

import structlog
import httpx

logger = structlog.get_logger()


class FileParser:
    """
    File parser for extracting text from uploaded documents.
    Supports: PDF (via pdfplumber), DOCX (via python-docx), plain text.
    """

    async def extract_from_url(self, file_url: str) -> str:
        """
        Download file from URL and extract text content.

        Args:
            file_url: URL to download the file from

        Returns:
            Extracted text content
        """
        logger.info("file_parser.download.start", url=file_url[:100])

        try:
            async with httpx.AsyncClient(timeout=60.0) as client:
                response = await client.get(file_url)
                response.raise_for_status()
                content = response.content

            content_type = response.headers.get("content-type", "")
            file_url_lower = file_url.lower()

            if file_url_lower.endswith(".pdf") or "pdf" in content_type:
                text = self._parse_pdf(content)
            elif file_url_lower.endswith(".docx") or "wordprocessingml" in content_type:
                text = self._parse_docx(content)
            elif file_url_lower.endswith(".doc"):
                # .doc format not natively supported, try as text
                logger.warning("file_parser.unsupported_format", format=".doc")
                text = content.decode("utf-8", errors="replace")
            elif file_url_lower.endswith(".txt") or "text/" in content_type:
                text = content.decode("utf-8", errors="replace")
            else:
                # Attempt to decode as text
                text = content.decode("utf-8", errors="replace")

            logger.info(
                "file_parser.complete",
                url=file_url[:100],
                text_length=len(text),
            )
            return text

        except httpx.HTTPStatusError as e:
            logger.error(
                "file_parser.http_error",
                url=file_url[:100],
                status=e.response.status_code,
            )
            raise ValueError(f"Failed to download file: HTTP {e.response.status_code}")
        except Exception as e:
            logger.error(
                "file_parser.error",
                url=file_url[:100],
                error=str(e),
            )
            raise

    def extract_from_bytes(self, content: bytes, filename: str) -> str:
        """
        Extract text from file bytes.

        Args:
            content: Raw file bytes
            filename: Original filename (used to detect format)

        Returns:
            Extracted text content
        """
        filename_lower = filename.lower()

        if filename_lower.endswith(".pdf"):
            return self._parse_pdf(content)
        elif filename_lower.endswith(".docx"):
            return self._parse_docx(content)
        elif filename_lower.endswith(".txt"):
            return content.decode("utf-8", errors="replace")
        else:
            return content.decode("utf-8", errors="replace")

    def _parse_pdf(self, content: bytes) -> str:
        """Extract text from PDF using pdfplumber."""
        try:
            import pdfplumber
        except ImportError:
            logger.error("file_parser.pdf.import_error", error="pdfplumber not installed")
            raise RuntimeError("pdfplumber is required for PDF parsing")

        pages_text = []
        page_count = 0

        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                page_count = len(pdf.pages)
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)

            result = "\n\n".join(pages_text)
            logger.info(
                "file_parser.pdf.complete",
                page_count=page_count,
                text_length=len(result),
            )
            return result
        except Exception as e:
            logger.error("file_parser.pdf.error", error=str(e))
            raise ValueError(f"Failed to parse PDF: {str(e)}")

    def _parse_docx(self, content: bytes) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document
        except ImportError:
            logger.error("file_parser.docx.import_error", error="python-docx not installed")
            raise RuntimeError("python-docx is required for DOCX parsing")

        try:
            doc = Document(io.BytesIO(content))
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            result = "\n\n".join(paragraphs)
            logger.info(
                "file_parser.docx.complete",
                paragraphs=len(paragraphs),
                text_length=len(result),
            )
            return result
        except Exception as e:
            logger.error("file_parser.docx.error", error=str(e))
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

    def get_page_count(self, content: bytes, filename: str) -> int:
        """Get page count for a file (PDF only)."""
        if filename.lower().endswith(".pdf"):
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(content)) as pdf:
                    return len(pdf.pages)
            except Exception:
                return 0
        return 0


# Singleton instance
file_parser = FileParser()
