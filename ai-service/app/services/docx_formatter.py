"""DOCX export helpers for Vietnamese lesson plans.

The formatter intentionally targets the formal CV-5512 lesson-plan shape:
plain Word document typography, Times New Roman, Letter page geometry, and
full-width teacher/student activity tables.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

import structlog

logger = structlog.get_logger()

try:  # Keep this module importable even when python-docx is not installed.
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Twips
except ImportError:  # pragma: no cover - exercised by runtime environments.
    Document = None  # type: ignore[assignment]
    WD_CELL_VERTICAL_ALIGNMENT = None  # type: ignore[assignment]
    WD_TABLE_ALIGNMENT = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    OxmlElement = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    Twips = None  # type: ignore[assignment]


FONT_NAME = "Times New Roman"
BODY_FONT_PT = 13
PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
PAGE_MARGIN_DXA = 1440
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - (PAGE_MARGIN_DXA * 2)
TABLE_CELL_MARGIN_DXA = 120

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BOLD_LABEL_RE = re.compile(r"^\*\*(?P<label>[^*]+?)\s*:\*\*\s*(?P<value>.*)$")
_PLAIN_LABEL_RE = re.compile(r"^(?P<label>[^:]{1,40})\s*:\s*(?P<value>.+)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.+?)\s*$")
_ORDERED_RE = re.compile(r"^\s*\d+[.)]\s+(.+?)\s*$")
_MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


@dataclass
class MarkdownTable:
    headers: list[str]
    rows: list[list[str]]


def generate_docx_bytes(plan: dict, markdown: str) -> bytes:
    """Generate CV-5512-style DOCX bytes from final markdown or plan fallback."""
    if Document is None:
        logger.warning("formatter.docx.import_error", error="python-docx not installed")
        return b""

    try:
        doc = Document()
        _apply_document_layout(doc)
        source_markdown = markdown.strip() or _markdown_from_plan(plan)
        _render_markdown_document(doc, source_markdown, plan)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
    except Exception as exc:
        logger.error("formatter.docx.generate_error", error=str(exc), exc_info=True)
        return b""


def _apply_document_layout(doc: Any) -> None:
    section = doc.sections[0]
    section.page_width = Twips(PAGE_WIDTH_DXA)
    section.page_height = Twips(PAGE_HEIGHT_DXA)
    section.top_margin = Twips(PAGE_MARGIN_DXA)
    section.right_margin = Twips(PAGE_MARGIN_DXA)
    section.bottom_margin = Twips(PAGE_MARGIN_DXA)
    section.left_margin = Twips(PAGE_MARGIN_DXA)
    section.header_distance = Twips(720)
    section.footer_distance = Twips(720)

    _configure_style(doc, "Normal", bold=False, size_pt=BODY_FONT_PT)
    _configure_style(doc, "Title", bold=True, size_pt=BODY_FONT_PT)
    _configure_style(doc, "Heading 1", bold=True, size_pt=BODY_FONT_PT)
    _configure_style(doc, "Heading 2", bold=True, size_pt=BODY_FONT_PT)
    _configure_style(doc, "Heading 3", bold=True, size_pt=BODY_FONT_PT)
    _configure_style(doc, "List Bullet", bold=False, size_pt=BODY_FONT_PT)
    _configure_style(doc, "List Number", bold=False, size_pt=BODY_FONT_PT)


def _configure_style(doc: Any, style_name: str, *, bold: bool, size_pt: int) -> None:
    try:
        style = doc.styles[style_name]
    except KeyError:
        return

    style.font.name = FONT_NAME
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    _set_element_fonts(style.element)
    paragraph_format = getattr(style, "paragraph_format", None)
    if paragraph_format is not None:
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing = 1.15


def _set_element_fonts(element: Any) -> None:
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT_NAME)


def _render_markdown_document(doc: Any, markdown: str, plan: dict) -> None:
    lines = [line.rstrip() for line in markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    start_index = _render_opening_block(doc, lines, plan)

    index = start_index
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.strip()
        if not line:
            index += 1
            continue

        parsed_table, next_index, malformed_lines = _parse_table(lines, index)
        if parsed_table is not None:
            _add_table(doc, parsed_table)
            index = next_index
            continue
        if malformed_lines:
            _add_malformed_table_lines(doc, malformed_lines)
            index = next_index
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            _add_heading(doc, _clean_inline_text(heading.group(2)), level)
            index += 1
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            _add_list_item(doc, bullet.group(1), style="List Bullet")
            index += 1
            continue

        ordered = _ORDERED_RE.match(line)
        if ordered:
            _add_list_item(doc, ordered.group(1), style="List Number")
            index += 1
            continue

        _add_body_paragraph(doc, line)
        index += 1


def _render_opening_block(doc: Any, lines: list[str], plan: dict) -> int:
    metadata = dict(_metadata_from_plan(plan))
    title = "KẾ HOẠCH BÀI DẠY"
    extra_lines: list[str] = []
    index = 0

    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            break

        heading = _HEADING_RE.match(line)
        if heading and len(heading.group(1)) == 1:
            title = _clean_inline_text(heading.group(2))
            index += 1
            continue

        label = _parse_metadata_line(line)
        if label:
            key, value = label
            metadata[key] = value
        else:
            extra_lines.append(line)
        index += 1

    _add_title_block(doc, title, metadata)
    for extra_line in extra_lines:
        _add_body_paragraph(doc, extra_line)
    return index


def _metadata_from_plan(plan: dict) -> dict[str, str]:
    metadata = plan.get("metadata", {}) if isinstance(plan, dict) else {}
    result: dict[str, str] = {}
    if metadata.get("subject"):
        result["subject"] = str(metadata["subject"])
    if metadata.get("grade"):
        result["grade"] = str(metadata["grade"])
    if metadata.get("topic"):
        result["topic"] = str(metadata["topic"])
    if metadata.get("duration_minutes"):
        result["duration"] = f"{metadata['duration_minutes']} phút"
    return result


def _parse_metadata_line(line: str) -> tuple[str, str] | None:
    match = _BOLD_LABEL_RE.match(line) or _PLAIN_LABEL_RE.match(_strip_markdown_wrappers(line))
    if not match:
        return None

    label = _clean_inline_text(match.group("label")).lower()
    value = _clean_inline_text(match.group("value"))
    if not value:
        return None
    if "môn" in label:
        return "subject", value
    if "lớp" in label:
        return "grade", value
    if "bài" in label or "chủ đề" in label:
        return "topic", value
    if "thời" in label:
        return "duration", value
    return None


def _add_title_block(doc: Any, title: str, metadata: dict[str, str]) -> None:
    title_text = title.strip() or "KẾ HOẠCH BÀI DẠY"
    _add_paragraph(
        doc,
        title_text.upper(),
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        style="Title",
    )

    topic = metadata.get("topic", "").strip()
    if topic and topic.lower() != title_text.lower():
        _add_paragraph(
            doc,
            topic.upper(),
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            style="Title",
        )

    duration = metadata.get("duration", "").strip()
    if duration:
        normalized_duration = duration
        if "thời" not in normalized_duration.lower() and "phút" in normalized_duration.lower():
            normalized_duration = f"thời lượng: {normalized_duration}"
        _add_paragraph(
            doc,
            f"({normalized_duration})",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    subject = metadata.get("subject", "").strip()
    grade = metadata.get("grade", "").strip()
    if subject or grade:
        parts = []
        if subject:
            parts.append(f"Môn học: {subject}")
        if grade:
            parts.append(f"Lớp: {grade}")
        _add_paragraph(doc, " - ".join(parts), alignment=WD_ALIGN_PARAGRAPH.CENTER)


def _add_heading(doc: Any, text: str, level: int) -> None:
    if level <= 1:
        _add_paragraph(doc, text.upper(), bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, style="Title")
        return

    style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
    paragraph = _add_paragraph(doc, text, bold=True, style=style)
    paragraph.paragraph_format.keep_with_next = True


def _add_body_paragraph(doc: Any, text: str) -> Any:
    return _add_paragraph(doc, _strip_markdown_wrappers(text), alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)


def _add_list_item(doc: Any, text: str, *, style: str) -> Any:
    paragraph = doc.add_paragraph(style=style)
    _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _add_inline_runs(paragraph, _strip_markdown_wrappers(text))
    return paragraph


def _add_paragraph(
    doc: Any,
    text: str,
    *,
    bold: bool = False,
    alignment: Any | None = None,
    style: str | None = None,
) -> Any:
    paragraph = doc.add_paragraph(style=style)
    _format_paragraph(paragraph, alignment=alignment or WD_ALIGN_PARAGRAPH.JUSTIFY)
    _add_inline_runs(paragraph, text, default_bold=bold)
    return paragraph


def _format_paragraph(paragraph: Any, *, alignment: Any) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15


def _add_inline_runs(paragraph: Any, text: str, *, default_bold: bool = False) -> None:
    cleaned = _clean_inline_text(text, preserve_bold=True)
    parts = re.split(r"(\*\*.+?\*\*)", cleaned)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        display = part[2:-2] if is_bold else part
        run = paragraph.add_run(display)
        run.bold = default_bold or is_bold
        run.font.name = FONT_NAME
        run.font.size = Pt(BODY_FONT_PT)
        _set_run_fonts(run)


def _set_run_fonts(run: Any) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT_NAME)


def _parse_table(lines: list[str], start: int) -> tuple[MarkdownTable | None, int, list[str]]:
    line = lines[start].strip()
    if not _looks_like_table_line(line):
        return None, start, []

    block: list[str] = []
    index = start
    while index < len(lines) and _looks_like_table_line(lines[index].strip()):
        block.append(lines[index].strip())
        index += 1

    if len(block) < 2 or not _is_table_separator(block[1]):
        return None, index, block

    headers = _split_table_row(block[0])
    rows = [_split_table_row(row) for row in block[2:] if not _is_table_separator(row)]
    if not headers or any(len(row) != len(headers) for row in rows):
        return None, index, block

    return MarkdownTable(headers=headers, rows=rows), index, []


def _looks_like_table_line(line: str) -> bool:
    return line.startswith("|") and line.endswith("|") and line.count("|") >= 2


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def _split_table_row(line: str) -> list[str]:
    return [_clean_inline_text(cell.strip()) for cell in line.strip().strip("|").split("|")]


def _add_table(doc: Any, markdown_table: MarkdownTable) -> None:
    column_count = len(markdown_table.headers)
    table = doc.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    try:
        table.style = "Table Grid"
    except KeyError:
        pass

    widths = _table_column_widths(column_count)
    _apply_table_geometry(table, widths)

    header_cells = table.rows[0].cells
    for column, text in enumerate(markdown_table.headers):
        _write_cell(header_cells[column], text, bold=True)
    _mark_header_row(table.rows[0])

    for row in markdown_table.rows:
        cells = table.add_row().cells
        for column, text in enumerate(row):
            _write_cell(cells[column], text, bold=False)
    _apply_table_geometry(table, widths)


def _table_column_widths(column_count: int) -> list[int]:
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def _apply_table_geometry(table: Any, widths: list[int]) -> None:
    table_width = sum(widths)
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(table_width))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is not None:
        table._tbl.remove(tbl_grid)
    tbl_grid = OxmlElement("w:tblGrid")
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    table._tbl.insert(1, tbl_grid)

    for row in table.rows:
        for column, cell in enumerate(row.cells):
            _apply_cell_geometry(cell, widths[column])


def _apply_cell_geometry(cell: Any, width: int) -> None:
    cell.width = Twips(width)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width))

    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        margin = tc_mar.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), str(TABLE_CELL_MARGIN_DXA))
        margin.set(qn("w:type"), "dxa")


def _write_cell(cell: Any, text: str, *, bold: bool) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _add_inline_runs(paragraph, text, default_bold=bold)


def _mark_header_row(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _add_malformed_table_lines(doc: Any, lines: list[str]) -> None:
    for line in lines:
        if _is_table_separator(line):
            continue
        cells = [cell for cell in _split_table_row(line) if cell]
        if cells:
            _add_body_paragraph(doc, " | ".join(cells))


def _clean_inline_text(text: str, *, preserve_bold: bool = False) -> str:
    cleaned = text.replace("`", "").replace("\\", "")
    cleaned = cleaned.replace("<br>", " ").replace("<br/>", " ").replace("<br />", " ")
    cleaned = _MARKDOWN_LINK_RE.sub(r"\1", cleaned)
    cleaned = cleaned.strip()
    if preserve_bold:
        return cleaned
    return cleaned.replace("**", "").replace("__", "").strip()


def _strip_markdown_wrappers(text: str) -> str:
    return text.strip().strip("`").strip()


def _markdown_from_plan(plan: dict) -> str:
    metadata = plan.get("metadata", {}) if isinstance(plan, dict) else {}
    sections = plan.get("sections", {}) if isinstance(plan, dict) else {}
    lines = [
        "# KẾ HOẠCH BÀI DẠY",
        f"**Môn học:** {metadata.get('subject', '')}",
        f"**Lớp:** {metadata.get('grade', '')}",
        f"**Bài:** {metadata.get('topic', '')}",
        f"**Thời gian:** {metadata.get('duration_minutes', 45)} phút",
        "",
        "## I. MỤC TIÊU",
        "### 1. Năng lực",
    ]
    for item in metadata.get("competencies") or metadata.get("objectives") or []:
        lines.append(f"- {item}")
    lines.extend(["### 2. Phẩm chất"])
    for item in metadata.get("qualities") or []:
        lines.append(f"- {item}")

    lines.extend([
        "",
        "## II. PHƯƠNG PHÁP DẠY HỌC VÀ KĨ THUẬT DẠY HỌC",
    ])
    for item in metadata.get("methods") or []:
        lines.append(f"- {item}")

    lines.extend([
        "",
        "## III. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU",
    ])
    for item in metadata.get("materials") or []:
        lines.append(f"- {item}")

    lines.extend(["", "## IV. TIẾN TRÌNH DẠY HỌC"])
    for index, (key, section) in enumerate(sections.items(), 1):
        if not isinstance(section, dict):
            continue
        title = section.get("title") or key.replace("_", " ").title()
        lines.extend([
            f"### Hoạt động {index}: {title}",
            _field_line("Thời lượng", section.get("duration")),
            _field_line("Mục tiêu", section.get("objective")),
            _field_line("Nội dung", section.get("content")),
            _field_line("Sản phẩm", section.get("product")),
            _field_line("Tổ chức thực hiện", section.get("organization")),
            _field_line("Phương án đánh giá", section.get("assessment")),
            "",
        ])

    return "\n".join(line for line in lines if line is not None)


def _field_line(label: str, value: Any) -> str | None:
    if value is None or value == "":
        return None
    if label == "Thời lượng" and isinstance(value, int):
        return f"{label}: {value} phút"
    return f"**{label}:** {value}"
