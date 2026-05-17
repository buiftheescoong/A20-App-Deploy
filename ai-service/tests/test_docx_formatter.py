from io import BytesIO

from docx import Document

from app.graph.nodes.formatter import _generate_docx_bytes
from tests.test_pipeline import _quality_plan_fixture


def _open_docx(markdown: str, plan: dict | None = None):
    default_plan = {
        "metadata": {
            "subject": "Toán",
            "grade": "8",
            "topic": "Hàm số bậc nhất",
            "teaching_model": "CV-5512",
            "duration_minutes": 45,
        },
        "sections": {},
    }
    docx_bytes = _generate_docx_bytes(plan or default_plan, markdown)
    assert docx_bytes
    return Document(BytesIO(docx_bytes))


def _paragraph_texts(doc) -> list[str]:
    return [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]


def test_docx_formatter_uses_cv5512_page_geometry():
    doc = _open_docx(_quality_plan_fixture("CV-5512"))
    section = doc.sections[0]

    assert section.page_width.twips == 12240
    assert section.page_height.twips == 15840
    assert section.top_margin.twips == 1440
    assert section.right_margin.twips == 1440
    assert section.bottom_margin.twips == 1440
    assert section.left_margin.twips == 1440


def test_docx_formatter_keeps_title_and_section_order():
    doc = _open_docx(_quality_plan_fixture("CV-5512"))
    texts = _paragraph_texts(doc)

    assert texts[0] == "KẾ HOẠCH BÀI DẠY"
    assert "HÀM SỐ BẬC NHẤT" in texts[1]

    section_titles = [
        "I. MỤC TIÊU",
        "II. PHƯƠNG PHÁP DẠY HỌC VÀ KĨ THUẬT DẠY HỌC",
        "III. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU",
        "IV. TIẾN TRÌNH DẠY HỌC",
    ]
    indexes = [texts.index(title) for title in section_titles]
    assert indexes == sorted(indexes)


def test_docx_formatter_renders_activity_tables_as_real_word_tables():
    doc = _open_docx(_quality_plan_fixture("CV-5512"))

    assert len(doc.tables) == 4
    table = doc.tables[0]
    assert len(table.columns) == 2
    assert len(table.rows) == 5
    assert "Hoạt động của giáo viên" in table.cell(0, 0).text
    assert "Hoạt động của học sinh" in table.cell(0, 1).text

    first_column_text = "\n".join(row.cells[0].text for row in table.rows)
    assert "Chuyển giao nhiệm vụ học tập" in first_column_text
    assert "Thực hiện nhiệm vụ học tập" in first_column_text
    assert "Báo cáo thảo luận" in first_column_text
    assert "Kết luận, nhận định" in first_column_text


def test_docx_formatter_preserves_teacher_student_table_content():
    markdown = _quality_plan_fixture("CV-5512").replace(
        "Chuyển giao nhiệm vụ học tập: nêu nhiệm vụ",
        "Chuyển giao nhiệm vụ học tập: giao phiếu phân tích số liệu",
    ).replace(
        "Tiếp nhận nhiệm vụ",
        "Tiếp nhận phiếu và phân công nhóm",
        1,
    )

    doc = _open_docx(markdown)
    table = doc.tables[0]

    assert "giao phiếu phân tích số liệu" in table.cell(1, 0).text
    assert "Tiếp nhận phiếu và phân công nhóm" in table.cell(1, 1).text


def test_docx_formatter_falls_back_to_readable_text_for_malformed_tables():
    markdown = """
# KẾ HOẠCH BÀI DẠY
**Bài:** Bài kiểm thử
**Thời gian:** 45 phút

## IV. TIẾN TRÌNH DẠY HỌC
| Hoạt động của giáo viên | Hoạt động của học sinh |
| Chuyển giao nhiệm vụ học tập | Học sinh tiếp nhận |
"""

    doc = _open_docx(markdown)
    texts = _paragraph_texts(doc)

    assert len(doc.tables) == 0
    assert any("Hoạt động của giáo viên | Hoạt động của học sinh" in text for text in texts)
    assert any("Chuyển giao nhiệm vụ học tập | Học sinh tiếp nhận" in text for text in texts)


def test_docx_formatter_uses_plan_fallback_when_markdown_empty():
    plan = {
        "metadata": {
            "subject": "Sinh học",
            "grade": "8",
            "topic": "Trao đổi chất",
            "teaching_model": "CV-5512",
            "duration_minutes": 45,
            "competencies": ["Mô tả được vai trò của trao đổi chất."],
            "qualities": ["Trách nhiệm khi làm việc nhóm."],
            "materials": ["SGK Sinh học 8", "Phiếu học tập"],
        },
        "sections": {
            "khoi_dong": {
                "title": "Khởi động",
                "duration": 5,
                "objective": "Kích hoạt hiểu biết ban đầu.",
                "content": "Quan sát tình huống mở đầu.",
                "product": "Câu trả lời nhanh.",
                "organization": "Giáo viên nêu câu hỏi, học sinh trả lời.",
                "assessment": "Quan sát câu trả lời.",
            }
        },
    }

    doc = _open_docx("", plan)
    texts = _paragraph_texts(doc)

    assert "TRAO ĐỔI CHẤT" in texts[1]
    assert "I. MỤC TIÊU" in texts
    assert any("Mô tả được vai trò của trao đổi chất." in text for text in texts)
    assert any("Hoạt động 1: Khởi động" in text for text in texts)
