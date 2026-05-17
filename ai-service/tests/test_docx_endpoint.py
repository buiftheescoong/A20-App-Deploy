from io import BytesIO

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.docx import router as docx_router
from tests.test_pipeline import _quality_plan_fixture


def test_docx_endpoint_returns_word_document_bytes():
    app = FastAPI()
    app.include_router(docx_router)
    client = TestClient(app)

    response = client.post(
        "/ai/docx",
        json={
            "plan_id": "plan-docx-endpoint",
            "plan": {
                "metadata": {
                    "subject": "Math",
                    "grade": "8",
                    "topic": "Linear functions",
                    "teaching_model": "CV-5512",
                    "duration_minutes": 45,
                },
                "sections": {},
            },
            "markdown": _quality_plan_fixture("CV-5512"),
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.content

    doc = Document(BytesIO(response.content))
    texts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    assert texts
    assert len(doc.tables) == 4
