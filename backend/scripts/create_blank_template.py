from docx import Document
from docx.shared import Pt
import os


def create_blank_gdpt_template(filename="blank_template.docx"):
    doc = Document()

    # Title
    title = doc.add_heading("KẾ HOẠCH BÀI DẠY (GIÁO ÁN) - GDPT 2018", 0)
    title.alignment = 1  # Center

    # Metadata
    doc.add_paragraph("Môn học: ....................")
    doc.add_paragraph("Lớp: ....................")
    doc.add_paragraph("Tên bài: ....................")
    doc.add_paragraph("Thời lượng: ....................")

    # I. Mục tiêu
    doc.add_heading("I. MỤC TIÊU BÀI HỌC", level=1)
    doc.add_paragraph("1. Về năng lực:")
    doc.add_paragraph("   - Thuộc bài học:")
    doc.add_paragraph("   - Năng lực chung:")
    doc.add_paragraph("2. Về phẩm chất:")

    # II. Thiết bị và học liệu
    doc.add_heading("II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU", level=1)
    doc.add_paragraph("1. Giáo viên: ....................")
    doc.add_paragraph("2. Học sinh: ....................")

    # III. Tiến trình
    doc.add_heading("III. TIẾN TRÌNH DẠY HỌC", level=1)

    # Hoạt động 1
    doc.add_heading("Hoạt động 1: Khởi động / Giao nhiệm vụ (Engage)", level=2)
    doc.add_paragraph("a) Mục tiêu: ....................")
    doc.add_paragraph("b) Nội dung: ....................")
    doc.add_paragraph("c) Sản phẩm: ....................")
    doc.add_paragraph("d) Tổ chức thực hiện: ....................")

    # Hoạt động 2
    doc.add_heading(
        "Hoạt động 2: Hình thành kiến thức mới / Khám phá (Explore)", level=2
    )
    doc.add_paragraph("a) Mục tiêu: ....................")
    doc.add_paragraph("b) Nội dung: ....................")
    doc.add_paragraph("c) Sản phẩm: ....................")
    doc.add_paragraph("d) Tổ chức thực hiện: ....................")

    # Hoạt động 3
    doc.add_heading(
        "Hoạt động 3: Luyện tập / Giải thích (Explain & Elaborate)", level=2
    )
    doc.add_paragraph("a) Mục tiêu: ....................")
    doc.add_paragraph("b) Nội dung: ....................")
    doc.add_paragraph("c) Sản phẩm: ....................")
    doc.add_paragraph("d) Tổ chức thực hiện: ....................")

    # Hoạt động 4
    doc.add_heading("Hoạt động 4: Vận dụng / Đánh giá (Evaluate)", level=2)
    doc.add_paragraph("a) Mục tiêu: ....................")
    doc.add_paragraph("b) Nội dung: ....................")
    doc.add_paragraph("c) Sản phẩm: ....................")
    doc.add_paragraph("d) Tổ chức thực hiện: ....................")

    # Save
    doc.save(filename)
    print(f"✅ Generated template: {os.path.abspath(filename)}")


if __name__ == "__main__":
    create_blank_gdpt_template()
